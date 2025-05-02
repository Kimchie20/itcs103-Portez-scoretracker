import tkinter as tk
from tkinter import messagebox, ttk
from openpyxl import Workbook, load_workbook
import os


excel_file = "student_scores.xlsx"
if not os.path.exists(excel_file):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Student Name", "Score", "Result"])
    workbook.save(excel_file)


workbook = load_workbook(excel_file)
worksheet = workbook.active

def save_student_score():
    student_name = name_entry.get().strip()
    score_input = score_entry.get().strip()

    if not student_name:
        messagebox.showerror("Input Error", "Student name cannot be empty.")
        return

    try:
        score = int(score_input)
        result = "Pass" if score >= 75 else "Fail"
        worksheet.append([student_name, score, result])
        workbook.save(excel_file)
        messagebox.showinfo("Success", f"Saved: {student_name} - {score} - {result}")
        name_entry.delete(0, tk.END)
        score_entry.delete(0, tk.END)
    except ValueError:
        messagebox.showerror("Input Error", "Please enter a valid number for score.")

def display_all_records():
    records_window = tk.Toplevel(root_window)
    records_window.title("All Student Records")

    tree = ttk.Treeview(records_window, columns=("Name", "Score", "Result"), show="headings")
    tree.heading("Name", text="Student Name")
    tree.heading("Score", text="Score")
    tree.heading("Result", text="Result")
    tree.pack(padx=10, pady=10, fill=tk.BOTH, expand=True)


    def refresh_tree():
        for item in tree.get_children():
            tree.delete(item)
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            tree.insert("", tk.END, values=row)

    refresh_tree()

    
    def edit_selected():
        selected = tree.selection()
        if not selected:
            messagebox.showwarning("No Selection", "Please select a record to edit.")
            return

        selected_values = tree.item(selected[0], "values")
        edit_window = tk.Toplevel(records_window)
        edit_window.title("Edit Record")

        tk.Label(edit_window, text="Name:").grid(row=0, column=0, padx=5, pady=5)
        edit_name = tk.Entry(edit_window)
        edit_name.grid(row=0, column=1, padx=5, pady=5)
        edit_name.insert(0, selected_values[0])

        tk.Label(edit_window, text="Score:").grid(row=1, column=0, padx=5, pady=5)
        edit_score = tk.Entry(edit_window)
        edit_score.grid(row=1, column=1, padx=5, pady=5)
        edit_score.insert(0, selected_values[1])

        def update_record():
            try:
                new_score = int(edit_score.get())
                new_result = "Pass" if new_score >= 75 else "Fail"
                for row in worksheet.iter_rows(min_row=2):
                    if row[0].value == selected_values[0] and str(row[1].value) == str(selected_values[1]):
                        row[0].value = edit_name.get()
                        row[1].value = new_score
                        row[2].value = new_result
                        workbook.save(excel_file)
                        refresh_tree()
                        edit_window.destroy()
                        messagebox.showinfo("Updated", "Record updated successfully.")
                        break
            except ValueError:
                messagebox.showerror("Error", "Score must be a number.")

        tk.Button(edit_window, text="Update", command=update_record).grid(row=2, column=0, columnspan=2, pady=10)

    
    def delete_selected():
        selected = tree.selection()
        if not selected:
            messagebox.showwarning("No Selection", "Please select a record to delete.")
            return

        selected_values = tree.item(selected[0], "values")
        confirm = messagebox.askyesno("Confirm Delete", f"Delete record: {selected_values[0]}?")
        if confirm:
            for i, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
                if row[0].value == selected_values[0] and str(row[1].value) == str(selected_values[1]):
                    worksheet.delete_rows(i)
                    workbook.save(excel_file)
                    refresh_tree()
                    messagebox.showinfo("Deleted", "Record deleted successfully.")
                    break

    
    tk.Button(records_window, text="Edit Selected", command=edit_selected).pack(pady=5)
    tk.Button(records_window, text="Delete Selected", command=delete_selected).pack(pady=5)


root_window = tk.Tk()
root_window.title("Student Score Tracker")

tk.Label(root_window, text="Student Name:").grid(row=0, column=0, padx=5, pady=5)
name_entry = tk.Entry(root_window)
name_entry.grid(row=0, column=1, padx=5, pady=5)

tk.Label(root_window, text="Score:").grid(row=1, column=0, padx=5, pady=5)
score_entry = tk.Entry(root_window)
score_entry.grid(row=1, column=1, padx=5, pady=5)

save_button = tk.Button(root_window, text="Save Score", command=save_student_score)
save_button.grid(row=2, column=0, columnspan=2, pady=10)

show_button = tk.Button(root_window, text="Show Records", command=display_all_records)
show_button.grid(row=3, column=0, columnspan=2, pady=5)

root_window.mainloop()